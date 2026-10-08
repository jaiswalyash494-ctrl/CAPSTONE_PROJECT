"""
RetainIQ Final College Black Book / Project Report Builder
Generates Ivan_Dsouza_RetainIQ_Blackbook.docx with professional college formatting,
authentic project metrics, UML diagrams, pipeline graphs, and complete technical chapters.
Target: 55-57 pages in PDF (Strict requirement: 51-60 pages).
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets inner margins for a table cell (in twips: 20 twips = 1 pt)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def set_table_borders(table, color="B0B8C0", sz="4", val="single"):
    """Applies clean, subtle borders to a table."""
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(borders)

def add_header_footer(doc):
    """Configures headers and footers with page numbers."""
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.1)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = True
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("RetainIQ: Survival Churn Modeling & Personalized Treatment Ranking")
        hrun.font.name = "Times New Roman"
        hrun.font.size = Pt(8.5)
        hrun.font.italic = True
        hrun.font.color.rgb = RGBColor(120, 130, 140)
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("KES' Shroff College of Arts & Commerce  |  T.Y. B.Sc. Data Science  |  Page ")
        frun.font.name = "Times New Roman"
        frun.font.size = Pt(9)
        frun.font.color.rgb = RGBColor(100, 110, 120)
        
        # Add dynamic page field
        fld = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        fp._p.append(fld)

def format_paragraph(p, space_before=0, space_after=4, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def add_body_p(doc, text="", bold_prefix="", italic_prefix="", space_before=0, space_after=4, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    format_paragraph(p, space_before, space_after, line_spacing, align)
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = "Times New Roman"
        r_bold.font.size = Pt(11)
        r_bold.font.bold = True
        r_bold.font.color.rgb = RGBColor(20, 30, 45)
    if italic_prefix:
        r_it = p.add_run(italic_prefix)
        r_it.font.name = "Times New Roman"
        r_it.font.size = Pt(11)
        r_it.font.italic = True
        r_it.font.color.rgb = RGBColor(40, 50, 60)
    if text:
        r_text = p.add_run(text)
        r_text.font.name = "Times New Roman"
        r_text.font.size = Pt(11)
        r_text.font.color.rgb = RGBColor(30, 35, 45)
    return p

def add_heading_1(doc, text, space_before=12, space_after=6):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=space_before, space_after=space_after, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 30, 65)  # Navy
    return p

def add_heading_2(doc, text, space_before=8, space_after=4):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=space_before, space_after=space_after, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(20, 45, 85)
    return p

def add_heading_3(doc, text, space_before=6, space_after=3):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=space_before, space_after=space_after, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 55, 95)
    return p

def add_figure(doc, img_path, caption, width_in=5.4):
    """Inserts a centered image with a styled caption."""
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(3)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_in))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(8)
        run_cap = p_cap.add_run(caption)
        run_cap.font.name = "Times New Roman"
        run_cap.font.size = Pt(9.5)
        run_cap.font.italic = True
        run_cap.font.bold = True
        run_cap.font.color.rgb = RGBColor(40, 50, 70)
    else:
        print(f"Warning: Image not found: {img_path}")

def add_code_listing(doc, code_str, caption):
    """Inserts a styled code box."""
    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.space_before = Pt(6)
    p_cap.paragraph_format.space_after = Pt(2)
    rc = p_cap.add_run(caption)
    rc.font.name = "Times New Roman"
    rc.font.size = Pt(9.5)
    rc.font.bold = True
    rc.font.color.rgb = RGBColor(20, 35, 65)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F6F8FA")
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(2)
    cp.paragraph_format.space_after = Pt(2)
    cp.paragraph_format.line_spacing = 1.0
    crun = cp.add_run(code_str.strip())
    crun.font.name = "Consolas"
    crun.font.size = Pt(8.2)
    crun.font.color.rgb = RGBColor(30, 40, 50)
    
    # Bottom margin after code table
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Report helper modules initialized.")
